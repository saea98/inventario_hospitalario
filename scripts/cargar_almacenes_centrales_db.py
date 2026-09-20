#!/usr/bin/env python3
"""
Carga directa de almacenes centrales a PostgreSQL (mismas reglas que
inventario.almacenes_centrales_carga), sin subir el Excel al servidor.

Uso:
  PGHOST=... PGPORT=5432 PGDATABASE=inventario_bd PGUSER=postgres PGPASSWORD=... \
    python scripts/cargar_almacenes_centrales_db.py [--aplicar]
"""

from __future__ import annotations

import argparse
import os
import sys
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import psycopg2
import psycopg2.extras

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = ROOT / 'data_import' / 'almacenes_centrales_251125.xlsx'
NA_VALUES = {'', 'N/A', 'NA', 'NONE', 'NULL', '-'}
FILA_INICIO = 13


def vacio(v) -> bool:
    return v is None or not str(v).strip()


def txt(v) -> str:
    return '' if v is None else str(v).strip()


def es_na(v: str) -> bool:
    return txt(v).upper() in NA_VALUES


def leer_filas(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        hoja = next(s for s in wb.sheetnames if s.strip().lower().startswith('almacenes'))
        ws = wb[hoja]
        out = []
        for i, row in enumerate(ws.iter_rows(min_row=FILA_INICIO, values_only=True), FILA_INICIO):
            entidad, clue_s, clue_i, unidad, direccion = (
                txt(row[0] if len(row) > 0 else None),
                txt(row[1] if len(row) > 1 else None),
                txt(row[2] if len(row) > 2 else None),
                txt(row[3] if len(row) > 3 else None),
                txt(row[4] if len(row) > 4 else None),
            )
            if not entidad and not clue_s and not unidad:
                continue
            if not es_na(clue_s):
                clue = clue_s[:20]
            elif not es_na(clue_i):
                clue = clue_i[:20]
            else:
                clue = None
            if es_na(clue_i):
                ib = clue if (es_na(clue_s) and clue) else None
            else:
                ib = clue_i[:20]
            out.append({
                'linea': i,
                'entidad': entidad,
                'clue': clue,
                'ib_clue': ib,
                'unidad': unidad,
                'direccion': direccion,
                'codigo': f'NAC-{clue}' if clue else None,
                'nombre_almacen': f'{entidad} - {unidad}'[:150] if entidad and unidad else '',
                'denominacion': unidad[:200] if unidad else '',
            })
        return out
    finally:
        wb.close()


def connect():
    return psycopg2.connect(
        host=os.environ['PGHOST'],
        port=os.environ.get('PGPORT', '5432'),
        dbname=os.environ['PGDATABASE'],
        user=os.environ['PGUSER'],
        password=os.environ['PGPASSWORD'],
        connect_timeout=20,
    )


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--xlsx', type=Path, default=DEFAULT_XLSX)
    ap.add_argument('--aplicar', action='store_true', help='Escribe en BD (default: dry-run)')
    args = ap.parse_args()
    dry = not args.aplicar

    if not args.xlsx.exists():
        print(f'No existe {args.xlsx}', file=sys.stderr)
        sys.exit(1)

    filas = leer_filas(args.xlsx)
    print(f'Filas Excel: {len(filas)}')
    print(f'Modo: {"DRY-RUN" if dry else "APLICAR"}')
    print(f'Host: {os.environ.get("PGHOST")} DB: {os.environ.get("PGDATABASE")}')

    conn = connect()
    cur = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)

    cur.execute('SELECT id FROM inventario_tipoinstitucion ORDER BY id LIMIT 1')
    tipo = cur.fetchone()
    if not tipo:
        print('ERROR: no hay TipoInstitucion', file=sys.stderr)
        sys.exit(1)
    tipo_id = tipo['id']
    now = datetime.now(timezone.utc)

    stats = {
        'inst_crear': 0, 'inst_act': 0, 'inst_omit': 0,
        'alm_crear': 0, 'alm_act': 0, 'alm_omit': 0, 'err': 0,
    }

    try:
        for f in filas:
            clue = f['clue']
            if not clue or not f['unidad'] or not f['entidad']:
                print(f"[error] fila {f['linea']}: datos incompletos")
                stats['err'] += 1
                continue

            cur.execute(
                'SELECT id, denominacion, nombre, direccion, estado, ib_clue '
                'FROM inventario_institucion WHERE clue=%s',
                (clue,),
            )
            inst = cur.fetchone()

            if inst is None:
                if f['ib_clue']:
                    cur.execute(
                        'SELECT id FROM inventario_institucion WHERE ib_clue=%s',
                        (f['ib_clue'],),
                    )
                    if cur.fetchone():
                        print(f"[error] fila {f['linea']}: ib_clue {f['ib_clue']} ya usado")
                        stats['err'] += 1
                        continue
                print(
                    f"[crear_institucion] {clue} ib={f['ib_clue'] or '—'} "
                    f"{f['denominacion']!r} ({f['entidad']})"
                )
                stats['inst_crear'] += 1
                if not dry:
                    cur.execute(
                        '''
                        INSERT INTO inventario_institucion
                        (clue, ib_clue, denominacion, estado, nombre, alcaldia_id,
                         tipo_institucion_id, direccion, telefono, email, activo,
                         fecha_creacion, fecha_actualizacion)
                        VALUES (%s,%s,%s,%s,%s,NULL,%s,%s,NULL,NULL,TRUE,%s,%s)
                        RETURNING id
                        ''',
                        (
                            clue, f['ib_clue'], f['denominacion'], f['entidad'],
                            f['denominacion'], tipo_id, f['direccion'] or None,
                            now, now,
                        ),
                    )
                    inst = {'id': cur.fetchone()['id']}
            else:
                cambios = {}
                if vacio(inst['denominacion']) and f['denominacion']:
                    cambios['denominacion'] = f['denominacion']
                if vacio(inst['nombre']) and f['denominacion']:
                    cambios['nombre'] = f['denominacion']
                if vacio(inst['direccion']) and f['direccion']:
                    cambios['direccion'] = f['direccion']
                if vacio(inst['estado']) and f['entidad']:
                    cambios['estado'] = f['entidad']
                if vacio(inst['ib_clue']) and f['ib_clue']:
                    cur.execute(
                        'SELECT id FROM inventario_institucion WHERE ib_clue=%s AND id<>%s',
                        (f['ib_clue'], inst['id']),
                    )
                    if not cur.fetchone():
                        cambios['ib_clue'] = f['ib_clue']
                if cambios:
                    print(f"[actualizar_institucion] {clue}: {cambios}")
                    stats['inst_act'] += 1
                    if not dry:
                        sets = ', '.join(f'{k}=%s' for k in cambios)
                        cur.execute(
                            f'UPDATE inventario_institucion SET {sets}, fecha_actualizacion=%s WHERE id=%s',
                            (*cambios.values(), now, inst['id']),
                        )
                else:
                    stats['inst_omit'] += 1

                if dry:
                    # necesitamos id para el almacén
                    pass

            if dry and inst is None:
                print(f"[crear_almacen] {f['codigo']} {f['nombre_almacen']!r} (inst nueva {clue})")
                stats['alm_crear'] += 1
                continue

            if inst is None:
                continue

            codigo = f['codigo']
            cur.execute(
                'SELECT id, nombre, direccion FROM inventario_almacen WHERE codigo=%s',
                (codigo,),
            )
            alm = cur.fetchone()
            if alm is None:
                print(f"[crear_almacen] {codigo} {f['nombre_almacen']!r} inst={clue}")
                stats['alm_crear'] += 1
                if not dry:
                    cur.execute(
                        '''
                        INSERT INTO inventario_almacen
                        (institucion_id, nombre, codigo, direccion, activo)
                        VALUES (%s,%s,%s,%s,TRUE)
                        ''',
                        (inst['id'], f['nombre_almacen'], codigo, f['direccion'] or None),
                    )
            else:
                cambios_a = {}
                if vacio(alm['nombre']) and f['nombre_almacen']:
                    cambios_a['nombre'] = f['nombre_almacen']
                if vacio(alm['direccion']) and f['direccion']:
                    cambios_a['direccion'] = f['direccion']
                if cambios_a:
                    print(f"[actualizar_almacen] {codigo}: {cambios_a}")
                    stats['alm_act'] += 1
                    if not dry:
                        sets = ', '.join(f'{k}=%s' for k in cambios_a)
                        cur.execute(
                            f'UPDATE inventario_almacen SET {sets} WHERE id=%s',
                            (*cambios_a.values(), alm['id']),
                        )
                else:
                    stats['alm_omit'] += 1

        print('\n=== Resumen ===')
        for k, v in stats.items():
            print(f'  {k}: {v}')

        if dry:
            conn.rollback()
            print('\nDry-run OK (rollback). Para aplicar: añade --aplicar')
        else:
            conn.commit()
            print('\nCambios aplicados (commit).')
    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()
        conn.close()


if __name__ == '__main__':
    main()
