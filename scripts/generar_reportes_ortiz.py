#!/usr/bin/env python3
"""
Llena plantillas Dr. Ortiz con existencia neta:
- EXISTENCIAS (neto) = fisica - comprometido
- Si neto > 0: tiempo y fecha de última entrada VACÍOS
- Si neto = 0: última fecha de ENTRADA + días sin recibir desde esa fecha
"""
from __future__ import annotations

import os
import re
from datetime import date
from pathlib import Path

import openpyxl
import psycopg2
from openpyxl.styles import Font


def connect():
    return psycopg2.connect(
        host=os.environ['PGHOST'],
        port=os.environ.get('PGPORT', '5432'),
        dbname=os.environ['PGDATABASE'],
        user=os.environ['PGUSER'],
        password=os.environ['PGPASSWORD'],
        connect_timeout=20,
    )


def normalize_clave(raw) -> str:
    if raw is None:
        return ''
    return re.sub(r'\s+', '', str(raw).strip().upper())


def clave_variants(clave: str) -> list[str]:
    c = normalize_clave(clave)
    if not c:
        return []
    variants = {c}
    if c.endswith('.00'):
        variants.add(c[:-3])
    else:
        variants.add(c + '.00')
    variants.add(c.replace('.', ''))
    return list(variants)


def resolve_match(clave: str, by_exact: dict):
    nc = normalize_clave(clave)
    if nc in by_exact:
        return by_exact[nc]
    for v in clave_variants(nc):
        if v in by_exact:
            return by_exact[v]
        for k, val in by_exact.items():
            if k.rstrip('.00') == v.rstrip('.00'):
                return val
    return None


def load_claves_from_xlsx(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    rows = []
    header_row = 2
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=False), start=1):
        vals = [str(c.value).strip().upper() if c.value is not None else '' for c in row]
        if any(v.startswith('CLAVE') for v in vals):
            header_row = i
            break
    for i, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        clave = normalize_clave(row[0] if row else None)
        if not clave:
            continue
        desc = (row[1] if len(row) > 1 and row[1] else '') or ''
        rows.append((i, clave, str(desc)))
    wb.close()
    return rows, header_row


def fetch_fisica_y_reserva(conn, claves: list[str]):
    cur = conn.cursor()
    variants = set()
    for c in claves:
        variants.update(clave_variants(c))
    if not variants:
        return {}
    variants_list = list(variants)

    cur.execute(
        """
        SELECT p.clave_cnis,
               COALESCE(SUM(l.cantidad_disponible), 0)::bigint AS fisica
        FROM inventario_producto p
        LEFT JOIN inventario_lote l ON l.producto_id = p.id
        WHERE UPPER(REPLACE(p.clave_cnis, ' ', '')) = ANY(%s)
           OR UPPER(REPLACE(p.clave_cnis, ' ', '')) || '.00' = ANY(%s)
           OR regexp_replace(UPPER(REPLACE(p.clave_cnis, ' ', '')), '\\.00$', '') = ANY(%s)
        GROUP BY p.clave_cnis
        """,
        (variants_list, variants_list, variants_list),
    )
    fisica_map = {normalize_clave(k): int(v or 0) for k, v in cur.fetchall()}

    cur.execute(
        """
        SELECT p.clave_cnis,
               COALESCE(SUM(la.cantidad_asignada), 0)::bigint AS reservado
        FROM inventario_loteasignado la
        JOIN inventario_loteubicacion lu ON lu.id = la.lote_ubicacion_id
        JOIN inventario_lote l ON l.id = lu.lote_id
        JOIN inventario_producto p ON p.id = l.producto_id
        WHERE la.surtido = false
          AND (
                UPPER(REPLACE(p.clave_cnis, ' ', '')) = ANY(%s)
             OR UPPER(REPLACE(p.clave_cnis, ' ', '')) || '.00' = ANY(%s)
             OR regexp_replace(UPPER(REPLACE(p.clave_cnis, ' ', '')), '\\.00$', '') = ANY(%s)
          )
        GROUP BY p.clave_cnis
        """,
        (variants_list, variants_list, variants_list),
    )
    reserva_map = {normalize_clave(k): int(v or 0) for k, v in cur.fetchall()}
    cur.close()

    result = {}
    for c in claves:
        nc = normalize_clave(c)
        fisica = resolve_match(nc, fisica_map) or 0
        reservado = resolve_match(nc, reserva_map) or 0
        result[nc] = {
            'fisica': int(fisica),
            'reservado': int(reservado),
            'neto': max(int(fisica) - int(reservado), 0),
        }
    return result


def fetch_ultima_entrada(conn, claves: list[str]) -> dict[str, date | None]:
    """Última fecha de movimiento ENTRADA / TRANSFERENCIA_ENTRADA / AJUSTE_POSITIVO."""
    cur = conn.cursor()
    variants = set()
    for c in claves:
        variants.update(clave_variants(c))
    if not variants:
        return {}
    variants_list = list(variants)
    cur.execute(
        """
        SELECT p.clave_cnis, MAX(m.fecha_movimiento)::date AS ultima
        FROM inventario_movimientoinventario m
        JOIN inventario_lote l ON l.id = m.lote_id
        JOIN inventario_producto p ON p.id = l.producto_id
        WHERE m.tipo_movimiento IN ('ENTRADA', 'TRANSFERENCIA_ENTRADA', 'AJUSTE_POSITIVO')
          AND COALESCE(m.anulado, false) = false
          AND (
                UPPER(REPLACE(p.clave_cnis, ' ', '')) = ANY(%s)
             OR UPPER(REPLACE(p.clave_cnis, ' ', '')) || '.00' = ANY(%s)
             OR regexp_replace(UPPER(REPLACE(p.clave_cnis, ' ', '')), '\\.00$', '') = ANY(%s)
          )
        GROUP BY p.clave_cnis
        """,
        (variants_list, variants_list, variants_list),
    )
    by_exact = {normalize_clave(k): v for k, v in cur.fetchall()}
    cur.close()
    result = {}
    for c in claves:
        nc = normalize_clave(c)
        result[nc] = resolve_match(nc, by_exact)
    return result


def dias_sin_recibir(ultima: date | None, hoy: date) -> str:
    if ultima is None:
        return 'Sin registro de entrada'
    dias = max((hoy - ultima).days, 0)
    return f'{dias} días'


def fill_workbook(src: Path, dst: Path, stock: dict, ultimas: dict, hoy: date):
    wb = openpyxl.load_workbook(src)
    ws = wb.active

    header_row = 2
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=False), start=1):
        vals = [str(c.value).strip().upper() if c.value is not None else '' for c in row]
        if any(v.startswith('CLAVE') for v in vals):
            header_row = i
            break

    # A CLAVE | B DESC | C EXISTENCIAS NETO | D TIEMPO | E ULTIMA ENTRADA | F FISICA | G COMPROMETIDO
    headers = {
        3: 'EXISTENCIAS (DISPONIBLE NETO)',
        4: 'TIEMPO QUE TIENEN DE NO ENTREGAR',
        5: 'ÚLTIMA FECHA DE ENTRADA',
        6: 'EXISTENCIA FISICA',
        7: 'COMPROMETIDO (PEDIDOS EN TRÁMITE)',
    }
    for col, title in headers.items():
        cell = ws.cell(row=header_row, column=col)
        cell.value = title
        cell.font = Font(bold=True)

    for i, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=False), start=header_row + 1):
        clave = normalize_clave(row[0].value)
        if not clave:
            continue
        info = stock.get(clave) or {'fisica': 0, 'reservado': 0, 'neto': 0}
        neto = info['neto']
        ws.cell(row=i, column=3).value = neto
        ws.cell(row=i, column=6).value = info['fisica']
        ws.cell(row=i, column=7).value = info['reservado']

        if neto > 0:
            # Con disponibilidad: tiempo y fecha vacíos
            ws.cell(row=i, column=4).value = None
            ws.cell(row=i, column=5).value = None
        else:
            ultima = ultimas.get(clave)
            ws.cell(row=i, column=4).value = dias_sin_recibir(ultima, hoy)
            ws.cell(row=i, column=5).value = ultima.strftime('%d/%m/%Y') if ultima else None

    last = ws.max_row + 2
    con_neto = sum(1 for v in stock.values() if v['neto'] > 0)
    sin_neto = sum(1 for v in stock.values() if v['neto'] <= 0)
    ws.cell(row=last, column=1).value = 'RESUMEN'
    ws.cell(row=last, column=1).font = Font(bold=True)
    ws.cell(row=last + 1, column=1).value = f'Claves con disponible neto > 0: {con_neto}'
    ws.cell(row=last + 2, column=1).value = f'Claves sin disponible neto: {sin_neto}'
    ws.cell(row=last + 3, column=1).value = (
        'EXISTENCIAS = física − comprometido (reservas activas). '
        'Si hay disponible: tiempo y última entrada vacíos. '
        'Si no hay: tiempo = días desde última ENTRADA/TRANSFERENCIA_ENTRADA/AJUSTE_POSITIVO.'
    )
    ws.cell(row=last + 4, column=1).value = f'Fecha de corte (stock actual): {hoy.isoformat()}'

    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)
    wb.close()
    return con_neto, sin_neto


def main():
    hoy = date.today()
    files = [
        Path('/Users/Sjimenez/Downloads/ARCHIVO DR ORTIZ MAT CURACIÓN.xlsx'),
        Path('/Users/Sjimenez/Downloads/ARCHIVO DR ORTIZ MEDICAMENTOS.xlsx'),
    ]
    out_dir = Path('/Users/Sjimenez/Downloads/reportes_ortiz')
    out_dir.mkdir(parents=True, exist_ok=True)

    conn = connect()
    print('DB connected', flush=True)

    for src in files:
        if not src.exists():
            print('MISSING', src)
            continue
        rows, _ = load_claves_from_xlsx(src)
        claves = [c for _, c, _ in rows]
        print(f'{src.name}: {len(claves)} claves', flush=True)
        stock = fetch_fisica_y_reserva(conn, claves)
        sin_neto = [c for c in claves if stock.get(normalize_clave(c), {}).get('neto', 0) <= 0]
        print(f'  sin neto: {len(sin_neto)} — consultando últimas entradas…', flush=True)
        ultimas = fetch_ultima_entrada(conn, sin_neto) if sin_neto else {}

        out_name = src.stem + f'_EXISTENCIAS_NETO_{hoy.strftime("%Y%m%d")}.xlsx'
        dst = out_dir / out_name
        con_e, sin_e = fill_workbook(src, dst, stock, ultimas, hoy)
        print(f'  -> {dst} (con_neto={con_e}, sin_neto={sin_e})', flush=True)

    conn.close()
    print('DONE', flush=True)


if __name__ == '__main__':
    main()
