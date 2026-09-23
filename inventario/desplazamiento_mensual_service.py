"""
Complemento de reporte mensual de desplazamiento (programa por entidad)
con surtimientos reales del sistema (LoteAsignado surtido en el mes/año).

Optimizado para no provocar 502:
- Primero lee CLUES del Excel y limita la consulta a esas instituciones
- Agrega en PostgreSQL (Sum / StringAgg) en lugar de iterar ORM fila a fila
- Evita copiar estilos celda a celda en miles de filas adicionales
"""

from __future__ import annotations

from calendar import monthrange
from collections import defaultdict
from datetime import date, datetime
from io import BytesIO
from typing import Dict, List, Optional, Set, Tuple

from django.contrib.postgres.aggregates import StringAgg
from django.db.models import CharField, F, Max, Q, Sum, Value
from django.db.models.functions import Coalesce, NullIf, Trim, Upper
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .pedidos_models import LoteAsignado

Key = Tuple[str, str]  # (clues_norm, clave_norm)

MAX_FILAS_ADICIONALES = 5000

MESES_NOMBRE = (
    '', 'Enero', 'Febrero', 'Marzo', 'Abril', 'Mayo', 'Junio',
    'Julio', 'Agosto', 'Septiembre', 'Octubre', 'Noviembre', 'Diciembre',
)

_MESES_EN_NOMBRE = {
    'ENERO': 1, 'FEBRERO': 2, 'MARZO': 3, 'ABRIL': 4, 'MAYO': 5, 'JUNIO': 6,
    'JULIO': 7, 'AGOSTO': 8, 'SEPTIEMBRE': 9, 'SETIEMBRE': 9, 'OCTUBRE': 10,
    'NOVIEMBRE': 11, 'DICIEMBRE': 12,
}


def mes_sugerido_en_nombre(nombre: str) -> Optional[int]:
    """Si el nombre del archivo menciona un mes, lo regresa (1-12)."""
    up = norm(nombre)
    for token, num in _MESES_EN_NOMBRE.items():
        if token in up:
            return num
    return None


def norm(v) -> str:
    return str(v).strip().upper() if v is not None else ''


def to_num(v) -> float:
    if v is None or v == '':
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def rango_mes(anio: int, mes: int) -> Tuple[date, date]:
    ultimo = monthrange(anio, mes)[1]
    return date(anio, mes, 1), date(anio, mes, ultimo)


def obtener_surtimientos_mes(
    anio: int,
    mes: int,
    *,
    clues_filtro: Optional[Set[str]] = None,
) -> Dict[Key, dict]:
    """
    Agrega cantidad surtida por (CLUES, CLAVE) en SQL.
    Si clues_filtro viene, solo instituciones de ese conjunto (ib_clue o clue).
    Indexa cada resultado por ib_clue y por clue (mismo dict) para el match del Excel.
    """
    inicio, fin = rango_mes(anio, mes)
    tz = timezone.get_current_timezone()
    dt_ini = timezone.make_aware(datetime.combine(inicio, datetime.min.time()), tz)
    dt_fin = timezone.make_aware(datetime.combine(fin, datetime.max.time()), tz)

    path_ib = 'item_propuesta__propuesta__solicitud__institucion_solicitante__ib_clue'
    path_clue = 'item_propuesta__propuesta__solicitud__institucion_solicitante__clue'
    path_clave = 'item_propuesta__producto__clave_cnis'
    path_desc = 'item_propuesta__producto__descripcion'
    path_obs = 'item_propuesta__propuesta__solicitud__observaciones_solicitud'
    path_folio = 'item_propuesta__propuesta__solicitud__folio'

    qs = (
        LoteAsignado.objects.filter(surtido=True, cantidad_asignada__gt=0)
        .annotate(
            fecha_efectiva=Coalesce(
                'fecha_surtimiento',
                'item_propuesta__propuesta__fecha_surtimiento',
            )
        )
        .filter(fecha_efectiva__gte=dt_ini, fecha_efectiva__lte=dt_fin)
    )

    if clues_filtro is not None:
        clues_list = [c for c in clues_filtro if c]
        if not clues_list:
            return {}
        qs = qs.filter(
            Q(**{f'{path_ib}__in': clues_list}) | Q(**{f'{path_clue}__in': clues_list})
        )

    folio_expr = Coalesce(
        NullIf(Trim(F(path_obs)), Value('')),
        NullIf(Trim(F(path_folio)), Value('')),
        Value(''),
        output_field=CharField(),
    )

    rows = qs.values(
        ib_clue=Trim(Upper(F(path_ib))),
        clue_ssa=Trim(Upper(F(path_clue))),
        clave=Trim(Upper(F(path_clave))),
    ).annotate(
        total=Sum('cantidad_asignada'),
        descripcion=Max(path_desc),
        folios=StringAgg(folio_expr, delimiter='|', distinct=True),
    )

    agg: Dict[Key, dict] = {}
    for row in rows.iterator(chunk_size=2000):
        clave = norm(row.get('clave'))
        if not clave:
            continue
        ib_clue = norm(row.get('ib_clue'))
        clue_ssa = norm(row.get('clue_ssa'))
        cant = float(row.get('total') or 0)
        if cant <= 0:
            continue
        folios_raw = (row.get('folios') or '').strip()
        folios = [f for f in folios_raw.split('|') if f][:8]
        desc = (row.get('descripcion') or '')[:300]
        info = {
            'cantidad': cant,
            'folios': folios,
            'descripcion': desc,
            'categoria': '',
            'clues_preferida': ib_clue or clue_ssa,
        }
        if ib_clue:
            agg[(ib_clue, clave)] = info
        if clue_ssa and (clue_ssa, clave) not in agg:
            agg[(clue_ssa, clave)] = info

    return agg


def encontrar_hoja(wb, *nombres):
    """Busca hoja por nombre (case-insensitive)."""
    por_lower = {str(n).strip().lower(): n for n in wb.sheetnames}
    for nom in nombres:
        key = str(nom).strip().lower()
        if key in por_lower:
            return wb[por_lower[key]]
    return None


def detectar_columnas(ws, header_row: int) -> Dict[str, int]:
    """
    Mapea columnas por encabezado. Soporta formato hospitalario (15 cols)
    y formato 1er nivel (12 cols).
    """
    cols: Dict[str, int] = {}
    max_c = min(ws.max_column or 20, 30)
    for c in range(1, max_c + 1):
        h = norm(ws.cell(header_row, c).value)
        if not h:
            continue
        # Cantidades primero: sus títulos también contienen "UNIDAD MEDICA"
        if 'PROYECTADA' in h:
            cols['proyectada'] = c
        elif 'REAL FISICA' in h or 'REAL FÍSICA' in h or (
            'ENTREGADA' in h and 'PROYECCION' not in h and 'PROYECCIÓN' not in h
        ):
            cols['entregada'] = c
        elif h == 'ENTIDAD' or h.startswith('ENTIDAD'):
            cols.setdefault('entidad', c)
        elif h == 'CLUES':
            cols['clues'] = c
        elif 'CATEGORIA' in h:
            cols['categoria'] = c
        elif h == 'CLAVE' or (h.startswith('CLAVE') and 'UNIC' not in h):
            cols.setdefault('clave', c)
        elif 'DESCRIPCION' in h or 'DESCRIPCIÓN' in h:
            cols.setdefault('descripcion', c)
        elif h == 'UNIDAD MEDICA' or h == 'UNIDAD MÉDICA' or (
            ('UNIDAD MEDICA' in h or 'UNIDAD MÉDICA' in h)
            and 'PROYECTADA' not in h
            and 'ENTREGADA' not in h
            and 'PROYECCION' not in h
            and 'PROYECCIÓN' not in h
            and 'EXISTENCIA' not in h
            and 'CPM' not in h
        ):
            cols.setdefault('unidad', c)
        elif 'FECHA PROGRAMADA' in h or h == 'FECHA':
            cols.setdefault('fecha', c)
        elif 'OBSERVACION' in h:
            cols['observaciones'] = c
        elif h == 'TIPO DE CLAVE' or h == 'TIPO':
            cols.setdefault('tipo', c)
        elif h == 'INSUMO':
            cols['insumo'] = c
    return cols


def complementar_workbook(
    archivo,
    anio: int,
    mes: int,
    *,
    nombre_fuente: str = '',
) -> Tuple[BytesIO, dict]:
    """
    Llena el programa mensual con surtimientos del mes.
    Orden: Excel → CLUES → query acotada → escritura ligera.
    Soporta hoja Datos/DATOS y layouts hospitalario / 1er nivel.
    """
    wb = load_workbook(archivo, keep_links=False)

    ws = encontrar_hoja(wb, 'Datos', 'DATOS')
    if ws is None:
        raise ValueError("El archivo debe contener la hoja 'Datos' (o 'DATOS').")

    header_row = None
    for r in range(1, min(25, (ws.max_row or 1) + 1)):
        vals = [norm(ws.cell(r, c).value) for c in range(1, min((ws.max_column or 15) + 1, 20))]
        if 'CLUES' in vals and any(v == 'CLAVE' or v.startswith('CLAVE') for v in vals):
            header_row = r
            break
    if not header_row:
        raise ValueError("No se encontró el renglón de encabezados con CLUES y CLAVE en la hoja Datos.")

    cols = detectar_columnas(ws, header_row)
    required = ('clues', 'clave', 'entregada')
    missing = [k for k in required if k not in cols]
    if missing:
        raise ValueError(
            f"Faltan columnas en Datos: {', '.join(missing)}. "
            f"Encabezados detectados en fila {header_row}."
        )

    col_clues = cols['clues']
    col_clave = cols['clave']
    col_entregada = cols['entregada']
    col_proyectada = cols.get('proyectada')
    col_obs = cols.get('observaciones')
    col_entidad = cols.get('entidad', 1)
    col_categoria = cols.get('categoria')
    col_unidad = cols.get('unidad')
    col_desc = cols.get('descripcion')
    col_fecha = cols.get('fecha')
    col_tipo = cols.get('tipo')
    col_insumo = cols.get('insumo')

    data_start = header_row + 1
    meta_clues = {}
    programa: Dict[Key, List[int]] = {}
    clues_en_excel: Set[str] = set()

    # Lectura por filas (más rápido que cell() suelto en archivos grandes)
    max_col = max(cols.values())
    for row in ws.iter_rows(min_row=data_start, max_row=ws.max_row, max_col=max_col):
        r = row[0].row
        clues = norm(row[col_clues - 1].value)
        clave = norm(row[col_clave - 1].value)
        if not clues or not clave:
            continue
        if clues in ('CLUES', 'ENTIDAD') or clave == 'CLAVE':
            continue
        programa.setdefault((clues, clave), []).append(r)
        clues_en_excel.add(clues)
        if clues not in meta_clues:
            cat_val = ''
            if col_categoria:
                cat_val = row[col_categoria - 1].value or ''
            unidad_val = ''
            if col_unidad:
                unidad_val = row[col_unidad - 1].value or ''
            meta_clues[clues] = {
                'entidad': (row[col_entidad - 1].value if col_entidad else None) or 'CIUDAD DE MEXICO',
                'categoria': cat_val,
                'unidad': unidad_val,
            }

    extra_pairs = set()
    wse = encontrar_hoja(wb, 'Extraordinario')
    if wse is not None:
        for r in range(8, min((wse.max_row or 0) + 1, 5000)):
            clues = norm(wse.cell(r, 3).value)
            clave = norm(wse.cell(r, 4).value)
            if clues and clave and not clues.startswith('ENTIDAD'):
                extra_pairs.add((clues, clave))
                clues_en_excel.add(clues)

    agg = obtener_surtimientos_mes(anio, mes, clues_filtro=clues_en_excel)

    fill_entregado = PatternFill('solid', fgColor='C6EFCE')
    fill_parcial = PatternFill('solid', fgColor='FFF2CC')
    fill_sin_exist = PatternFill('solid', fgColor='D9E2F3')
    fill_adicional = PatternFill('solid', fgColor='FCE4D6')
    font_ok = Font(color='006100', bold=True)
    font_info = Font(color='1F4E79')
    font_adic = Font(color='C65911', bold=True)

    periodo_txt = f'{MESES_NOMBRE[mes]} {anio}'
    aviso_archivo = ''
    mes_en_archivo = mes_sugerido_en_nombre(nombre_fuente or '')
    if mes_en_archivo and mes_en_archivo != mes:
        aviso_archivo = (
            f'AVISO: el archivo base menciona {MESES_NOMBRE[mes_en_archivo]} pero el filtro '
            f'de surtimientos es {periodo_txt}. La columna FECHA PROGRAMADA DE ENTREGA '
            f'pertenece al Excel cargado; las piezas sí corresponden a {periodo_txt}.'
        )

    # En archivos muy grandes (1er nivel ~37k filas) pintar cada "sin entrega"
    # satura CPU/memoria y provoca 502; solo marcamos filas con entrega.
    modo_ligero = len(programa) > 8000

    n_entregado = n_parcial = n_sin_exist = 0
    piezas_programa = piezas_adicionales = 0.0
    n_adicionales = 0
    cat_stats = defaultdict(lambda: {'clues': set(), 'claves': 0, 'piezas': 0.0, 'adic': 0.0, 'prog': 0.0})
    clues_all: Set[str] = set()

    def _folios_txt(info: dict) -> str:
        folios = info.get('folios') or []
        txt = ', '.join(folios[:6])
        if len(folios) > 6:
            txt += f' (+{len(folios) - 6} más)'
        return txt

    def _categoria_fila(r: int) -> str:
        if col_categoria:
            return str(ws.cell(r, col_categoria).value or 'Sin categoría').strip() or 'Sin categoría'
        return '1er nivel / sin categoría gerencial'

    for key, rows in programa.items():
        info = agg.get(key)
        clues, _clave = key
        for r in rows:
            proyectada = to_num(ws.cell(r, col_proyectada).value) if col_proyectada else 0.0
            cell_l = ws.cell(r, col_entregada)
            cell_obs = ws.cell(r, col_obs) if col_obs else None
            cat = _categoria_fila(r)
            clues_all.add(clues)

            if info and info['cantidad'] > 0:
                cant = info['cantidad']
                cell_l.value = cant
                piezas_programa += cant
                folios = _folios_txt(info)
                if proyectada and cant + 0.001 < proyectada:
                    n_parcial += 1
                    cell_l.fill = fill_parcial
                    nota = (
                        f'Surtido en {periodo_txt} (fecha de surtimiento). '
                        f'Entrega parcial ({cant:g} de {proyectada:g} proyectadas). '
                        f'Resto sujeto a existencia/disponibilidad. Folios: {folios}'
                    )
                else:
                    n_entregado += 1
                    cell_l.fill = fill_entregado
                    cell_l.font = font_ok
                    nota = (
                        f'Surtido en {periodo_txt} (fecha de surtimiento). '
                        f'Entregado conforme a programa. Folios: {folios}'
                    )
                if cell_obs is not None:
                    cell_obs.value = nota
                cat_stats[cat]['prog'] += cant
                cat_stats[cat]['piezas'] += cant
            else:
                n_sin_exist += 1
                if not modo_ligero:
                    cell_l.value = 0
                    cell_l.fill = fill_sin_exist
                    cell_l.font = font_info
                    if cell_obs is not None:
                        cell_obs.value = (
                            f'Sin entrega en {periodo_txt} por falta de existencia o disponibilidad '
                            'en almacén central (no implica falta de gestión).'
                        )
            cat_stats[cat]['clues'].add(clues)
            cat_stats[cat]['claves'] += 1

    last_data_row = max((max(rows) for rows in programa.values()), default=header_row)

    info_en_programa = {id(agg[k]) for k in programa if k in agg}
    vistos = set()
    adicionales = []
    for key, info in agg.items():
        oid = id(info)
        if oid in vistos or oid in info_en_programa:
            continue
        vistos.add(oid)
        if info['cantidad'] <= 0:
            continue
        clues_pref = info.get('clues_preferida') or key[0]
        adicionales.append((clues_pref, key[1], info))

    adicionales.sort(key=lambda x: (x[0], x[1]))
    truncados = max(0, len(adicionales) - MAX_FILAS_ADICIONALES)
    if truncados:
        adicionales = adicionales[:MAX_FILAS_ADICIONALES]

    for clues, clave, info in adicionales:
        last_data_row += 1
        n_adicionales += 1
        piezas_adicionales += info['cantidad']
        meta = meta_clues.get(clues, {})
        en_extra = (clues, clave) in extra_pairs
        folios = _folios_txt(info)
        cat = meta.get('categoria') or 'Entrega adicional'

        ws.cell(last_data_row, col_entidad).value = meta.get('entidad') or 'CIUDAD DE MEXICO'
        ws.cell(last_data_row, col_clues).value = clues
        if col_categoria:
            ws.cell(last_data_row, col_categoria).value = cat
        ws.cell(last_data_row, col_clave).value = clave
        if col_tipo:
            ws.cell(last_data_row, col_tipo).value = 'Entrega adicional'
        if col_insumo:
            ws.cell(last_data_row, col_insumo).value = 'Atención a demanda / fuera de programa'
        if col_desc:
            ws.cell(last_data_row, col_desc).value = info.get('descripcion') or ''
        if col_unidad:
            ws.cell(last_data_row, col_unidad).value = meta.get('unidad') or ''
        if col_proyectada:
            ws.cell(last_data_row, col_proyectada).value = 0
        cell_l = ws.cell(last_data_row, col_entregada)
        cell_l.value = info['cantidad']
        cell_l.fill = fill_adicional
        cell_l.font = font_adic
        if col_fecha:
            ws.cell(last_data_row, col_fecha).value = 'ENTREGA ADICIONAL'
        obs = (
            f'ENTREGA ADICIONAL / fuera del programa de desplazamiento. '
            f'Surtido en {periodo_txt} (fecha de surtimiento). Folios: {folios}'
        )
        if en_extra:
            obs += ' (También en Extraordinario).'
        if col_obs:
            ws.cell(last_data_row, col_obs).value = obs

        clues_all.add(clues)
        cat_stats[cat]['clues'].add(clues)
        cat_stats[cat]['claves'] += 1
        cat_stats[cat]['piezas'] += info['cantidad']
        cat_stats[cat]['adic'] += info['cantidad']

    if modo_ligero:
        extra_aviso = (
            f' Archivo grande ({len(programa)} líneas de programa): solo se colorearon '
            f'líneas con entrega; las sin surtimiento se dejaron sin marcar para evitar timeout.'
        )
        aviso_archivo = (aviso_archivo + extra_aviso).strip()

    _reescribir_resumen(
        wb,
        anio=anio,
        mes=mes,
        nombre_fuente=nombre_fuente,
        n_programa=len(programa),
        n_entregado=n_entregado,
        n_parcial=n_parcial,
        n_sin_exist=n_sin_exist,
        n_adicionales=n_adicionales,
        piezas_programa=piezas_programa,
        piezas_adicionales=piezas_adicionales,
        cat_stats=cat_stats,
        clues_all=clues_all,
        fill_entregado=fill_entregado,
        fill_parcial=fill_parcial,
        fill_sin_exist=fill_sin_exist,
        fill_adicional=fill_adicional,
        adicionales_truncados=truncados,
        aviso_archivo=aviso_archivo,
    )

    for name in list(wb.sheetnames):
        if name.lower() in ('entregas fuera de programa', 'resumen cruce entregas'):
            del wb[name]

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    total = piezas_programa + piezas_adicionales
    return buf, {
        'anio': anio,
        'mes': mes,
        'pares_programa': len(programa),
        'entregado_ok': n_entregado,
        'entregado_parcial': n_parcial,
        'sin_existencia': n_sin_exist,
        'adicionales': n_adicionales,
        'piezas_programa': piezas_programa,
        'piezas_adicionales': piezas_adicionales,
        'piezas_total': total,
        'pares_surtimiento_bd': len({id(v) for v in agg.values()}),
        'clues_filtradas': len(clues_en_excel),
        'adicionales_truncados': truncados,
        'modo_ligero': modo_ligero,
        'columnas': cols,
    }


def _reescribir_resumen(
    wb,
    *,
    anio,
    mes,
    nombre_fuente,
    n_programa,
    n_entregado,
    n_parcial,
    n_sin_exist,
    n_adicionales,
    piezas_programa,
    piezas_adicionales,
    cat_stats,
    clues_all,
    fill_entregado,
    fill_parcial,
    fill_sin_exist,
    fill_adicional,
    adicionales_truncados=0,
    aviso_archivo='',
):
    wsr = encontrar_hoja(wb, 'Resumen', 'RESUMEN')
    if wsr is None:
        wsr = wb.create_sheet('Resumen', 0)
    else:
        # openpyxl no permite escribir value en MergedCell; hay que descombinar primero
        for merged_range in list(wsr.merged_cells.ranges):
            wsr.unmerge_cells(str(merged_range))
        for row in wsr.iter_rows(min_row=1, max_row=80, max_col=12):
            for cell in row:
                cell.value = None
                cell.fill = PatternFill()
                cell.font = Font()

    total = piezas_programa + piezas_adicionales
    cobertura = (n_entregado + n_parcial) / max(n_programa, 1) * 100
    pct_adic = piezas_adicionales / max(total, 1) * 100
    periodo_txt = f'{MESES_NOMBRE[mes]} {anio}'

    title_font = Font(bold=True, size=14, color='1F4E79')
    kpi_fill = PatternFill('solid', fgColor='1F4E79')
    kpi_font = Font(bold=True, color='FFFFFF', size=11)
    pos_fill = PatternFill('solid', fgColor='C6EFCE')
    adic_fill = PatternFill('solid', fgColor='FCE4D6')
    warn_fill = PatternFill('solid', fgColor='FFF2CC')

    wsr['A1'] = f'AVANCE DE ENTREGAS — {MESES_NOMBRE[mes].upper()} {anio}'
    wsr['A1'].font = title_font
    wsr['A2'] = (
        f'Complementado con surtimientos del sistema (fecha de surtimiento = {periodo_txt}) | '
        f'Archivo base: {nombre_fuente or "programa mensual"} | '
        f'Generado: {datetime.now().strftime("%d/%m/%Y %H:%M")}'
    )
    wsr['A2'].font = Font(italic=True, size=10)
    if aviso_archivo:
        wsr['A3'] = aviso_archivo
        wsr['A3'].font = Font(bold=True, size=10, color='9C5700')
        wsr['A3'].fill = warn_fill
        wsr['A3'].alignment = Alignment(wrap_text=True)
        wsr.merge_cells('A3:F3')
    wsr['A4'] = 'INDICADORES DE GESTIÓN'
    wsr['A4'].font = Font(bold=True, size=12, color='833C0C')

    kpis = [
        (6, 'TOTAL PIEZAS ENTREGADAS', f'{total:,.0f}', 'total'),
        (7, 'Piezas del programa (Datos)', f'{piezas_programa:,.0f}', ''),
        (8, 'Piezas ENTREGA ADICIONAL (fuera de listado base)', f'{piezas_adicionales:,.0f}', 'adic'),
        (9, '% del volumen que fue atención adicional', f'{pct_adic:.1f}%', 'adic'),
        (10, 'CLUES atendidas', f'{len(clues_all)}', ''),
        (11, 'Líneas de programa con entrega', f'{n_entregado + n_parcial} de {n_programa} ({cobertura:.1f}%)', ''),
        (12, 'Líneas sin entrega por existencia/disponibilidad', f'{n_sin_exist}', ''),
        (13, 'Líneas de entrega adicional agregadas', f'{n_adicionales}', 'adic'),
    ]
    for row, label, val, kind in kpis:
        wsr.cell(row, 1, label).fill = kpi_fill
        wsr.cell(row, 1).font = kpi_font
        wsr.cell(row, 2, val).font = Font(bold=True, size=12)
        if kind == 'total':
            wsr.cell(row, 2).fill = pos_fill
        elif kind == 'adic':
            wsr.cell(row, 2).fill = adic_fill

    nota = (
        f'Criterio de piezas: LoteAsignado surtido con fecha_surtimiento en {periodo_txt}. '
        'La columna FECHA PROGRAMADA DE ENTREGA del Excel base no se modifica (es del programa cargado, '
        'no la fecha real de surtimiento). Folios con “-09-” u otro número son nombre de folio, no el mes. '
        'Lectura: azul en Datos = sin entrega por existencia/disponibilidad (no es falta de operación). '
        'Naranja = entrega adicional (atención a demanda fuera del programa base).'
    )
    if adicionales_truncados:
        nota += f' Se limitaron las adicionales a {MAX_FILAS_ADICIONALES} filas ({adicionales_truncados} omitidas).'
    wsr['A15'] = nota
    wsr['A15'].alignment = Alignment(wrap_text=True)
    wsr.merge_cells('A15:F17')

    wsr['A19'] = 'RESUMEN POR CATEGORÍA GERENCIAL (incluye adicionales)'
    wsr['A19'].font = Font(bold=True, size=12)
    headers = ['CATEGORIA GERENCIAL', 'CLUES', 'CLAVES/LÍNEAS', 'PIEZAS ENTREGADAS', 'DE PROGRAMA', 'ADICIONALES']
    for i, h in enumerate(headers, 1):
        cell = wsr.cell(20, i, h)
        cell.fill = PatternFill('solid', fgColor='833C0C')
        cell.font = Font(bold=True, color='FFFFFF')

    row_i = 21
    total_l = total_p = total_pg = total_ad = 0
    for cat, st in sorted(cat_stats.items(), key=lambda x: -x[1]['piezas']):
        wsr.cell(row_i, 1, cat)
        wsr.cell(row_i, 2, len(st['clues']))
        wsr.cell(row_i, 3, st['claves'])
        wsr.cell(row_i, 4, st['piezas'])
        wsr.cell(row_i, 5, st['prog'])
        wsr.cell(row_i, 6, st['adic'])
        total_l += st['claves']
        total_p += st['piezas']
        total_pg += st['prog']
        total_ad += st['adic']
        row_i += 1

    wsr.cell(row_i, 1, 'TOTAL').font = Font(bold=True)
    wsr.cell(row_i, 2, len(clues_all)).font = Font(bold=True)
    wsr.cell(row_i, 3, total_l).font = Font(bold=True)
    wsr.cell(row_i, 4, total_p).font = Font(bold=True)
    wsr.cell(row_i, 5, total_pg).font = Font(bold=True)
    wsr.cell(row_i, 6, total_ad).font = Font(bold=True)
    for col in range(1, 7):
        wsr.cell(row_i, col).fill = pos_fill

    row_i += 3
    wsr.cell(row_i, 1, 'LEYENDA EN HOJA DATOS').font = Font(bold=True)
    wsr.cell(row_i + 1, 1, 'Verde — Entregado conforme a programa').fill = fill_entregado
    wsr.cell(row_i + 2, 1, 'Amarillo — Entrega parcial (resto por existencia)').fill = fill_parcial
    wsr.cell(row_i + 3, 1, 'Azul — Sin entrega por falta de existencia/disponibilidad').fill = fill_sin_exist
    wsr.cell(row_i + 4, 1, 'Naranja — ENTREGA ADICIONAL / fuera del programa base').fill = fill_adicional

    wsr.column_dimensions['A'].width = 58
    wsr.column_dimensions['B'].width = 14
    wsr.column_dimensions['C'].width = 16
    wsr.column_dimensions['D'].width = 20
    wsr.column_dimensions['E'].width = 16
    wsr.column_dimensions['F'].width = 14


def asegurar_menu_desplazamiento_mensual():
    """Crea/actualiza la opción de menú (idempotente)."""
    from django.contrib.auth.models import Group

    from .models import MenuItemRol

    item, _ = MenuItemRol.objects.update_or_create(
        menu_item='desplazamiento_mensual',
        defaults={
            'nombre_mostrado': 'Desplazamiento mensual',
            'icono': 'fas fa-truck-loading',
            'url_name': 'reportes:desplazamiento_mensual',
            'orden': 53,
            'es_submenu': False,
            'activo': True,
        },
    )
    roles = Group.objects.filter(
        name__in=['Supervisión', 'Administrador', 'Gestor de Inventario', 'Analista']
    )
    if roles.exists():
        item.roles_permitidos.set(roles)
    return item
